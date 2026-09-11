from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"C:\NodeStack\resales2\output\documents\nezavisimaya-ai-platforma-kp.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "1F4D78"
BLUE2 = "2E74B5"
NAVY = "0B2545"
INK = "20252B"
MUTED = "5F6B76"
LIGHT = "F2F4F7"
PALE = "E8EEF5"
GREEN = "176B4D"
GREEN_FILL = "EAF5F0"
AMBER = "7A5A00"
AMBER_FILL = "FFF7E3"
RED = "9B1C1C"
RED_FILL = "FCEEEF"
WHITE = "FFFFFF"
DXA = 9360


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_font(run, size=None, bold=None, color=INK, italic=None, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    assert sum(widths) == DXA
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(DXA))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths[idx]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_keep(paragraph, next_=False):
    pPr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext" if next_ else "w:keepLines")
    pPr.append(keep)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_font(run, 8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

# Standard business brief, compact named overrides for page margins and title.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.25)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE2, 14, 7),
    ("Heading 2", 12.5, BLUE2, 10, 5),
    ("Heading 3", 11.2, BLUE, 7, 3),
]:
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = rgb(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name in ["List Bullet", "List Number"]:
    st = styles[style_name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(10.25)
    st.paragraph_format.left_indent = Inches(0.38)
    st.paragraph_format.first_line_indent = Inches(-0.19)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.08

for custom in ["Lead", "Small Note", "Table Text", "Table Head"]:
    if custom not in styles:
        styles.add_style(custom, WD_STYLE_TYPE.PARAGRAPH)
styles["Lead"].font.name = "Arial"
styles["Lead"].font.size = Pt(11.5)
styles["Lead"].font.bold = True
styles["Lead"].font.color.rgb = rgb(NAVY)
styles["Lead"].paragraph_format.space_after = Pt(7)
styles["Lead"].paragraph_format.line_spacing = 1.12
styles["Small Note"].font.name = "Arial"
styles["Small Note"].font.size = Pt(8.5)
styles["Small Note"].font.color.rgb = rgb(MUTED)
styles["Small Note"].paragraph_format.space_after = Pt(4)
styles["Table Text"].font.name = "Arial"
styles["Table Text"].font.size = Pt(8.5)
styles["Table Text"].font.color.rgb = rgb(INK)
styles["Table Text"].paragraph_format.space_after = Pt(0)
styles["Table Text"].paragraph_format.line_spacing = 1.02
styles["Table Head"].font.name = "Arial"
styles["Table Head"].font.size = Pt(8.3)
styles["Table Head"].font.bold = True
styles["Table Head"].font.color.rgb = rgb(NAVY)
styles["Table Head"].paragraph_format.space_after = Pt(0)

# Running header/footer.
hp = sec.header.paragraphs[0]
hp.text = "КОНЦЕПЦИЯ НЕЗАВИСИМОЙ AI-ПЛАТФОРМЫ  |  ВНУТРЕННИЙ ДОКУМЕНТ"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
for r in hp.runs:
    set_font(r, 8.2, True, MUTED)
fp = sec.footer.paragraphs[0]
add_page_field(fp)


def add_title(text_value, size=22, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text_value)
    set_font(r, size, True, NAVY)
    set_keep(p, True)
    return p


def add_subtitle(text_value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text_value)
    set_font(r, 11.5, False, MUTED)
    return p


def add_para(text_value, style=None, bold_prefix=None, after=None):
    p = doc.add_paragraph(style=style)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text_value.startswith(bold_prefix):
        a, b = text_value.split(":", 1)
        r = p.add_run(a + ":")
        set_font(r, bold=True, color=INK)
        r = p.add_run(b)
        set_font(r, color=INK)
    else:
        r = p.add_run(text_value)
        set_font(r, 10.25, color=INK)
    return p


def add_bullet(text_value):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text_value)
    set_font(r, 10.25, color=INK)
    return p


def add_callout(label, text_value, fill=PALE, label_color=BLUE, keep=True):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1760, 7600])
    table.style = "Table Grid"
    set_cell_shading(table.cell(0, 0), fill)
    set_cell_shading(table.cell(0, 1), fill)
    p0 = table.cell(0, 0).paragraphs[0]
    p0.style = styles["Table Head"]
    r = p0.add_run(label.upper())
    set_font(r, 8.4, True, label_color)
    p1 = table.cell(0, 1).paragraphs[0]
    p1.style = styles["Table Text"]
    r = p1.add_run(text_value)
    set_font(r, 9.5, True, NAVY)
    if keep:
        set_keep(p0)
        set_keep(p1)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_section_band(number, title):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [900, 8460])
    table.style = "Table Grid"
    set_cell_shading(table.cell(0, 0), NAVY)
    set_cell_shading(table.cell(0, 1), PALE)
    p0 = table.cell(0, 0).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p0.add_run(number)
    set_font(r, 11, True, WHITE)
    p1 = table.cell(0, 1).paragraphs[0]
    r = p1.add_run(title)
    set_font(r, 12.5, True, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_simple_table(headers, rows, widths, font_size=8.4, header_fill=PALE):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_shading(t.cell(0, i), header_fill)
        p = t.cell(0, i).paragraphs[0]
        p.style = styles["Table Head"]
        r = p.add_run(h)
        set_font(r, font_size, True, NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.style = styles["Table Text"]
            r = p.add_run(str(value))
            set_font(r, font_size, False, INK)
    set_table_geometry(t, widths)
    set_repeat_table_header(t.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# Opening memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(3)
r = p.add_run("СТРАТЕГИЧЕСКАЯ ЗАПИСКА")
set_font(r, 9, True, BLUE2)
add_title("Независимая AI-платформа анализа заявок и формирования коммерческих предложений")
add_subtitle("Первый прикладной контур: извлечение требований, классификация решения, подготовка описания и визуальной обложки КП")

meta = doc.add_table(rows=4, cols=2)
meta.style = "Table Grid"
set_table_geometry(meta, [1800, 7560])
for i, (label, value) in enumerate([
    ("Статус", "Концепция для управленческого решения"),
    ("Предмет", "Создание самостоятельного корпоративного AI-контура"),
    ("Горизонт", "Пилот 5-7 месяцев; промышленное внедрение 8-12 месяцев"),
    ("Дата", "28 августа 2026 г."),
]):
    set_cell_shading(meta.cell(i, 0), LIGHT)
    p0 = meta.cell(i, 0).paragraphs[0]
    r0 = p0.add_run(label)
    set_font(r0, 8.8, True, MUTED)
    p1 = meta.cell(i, 1).paragraphs[0]
    r1 = p1.add_run(value)
    set_font(r1, 9.2, i == 1, NAVY if i == 1 else INK)
doc.add_paragraph().paragraph_format.space_after = Pt(0)

add_callout(
    "Резюме",
    "Предлагается не отдельная функция PDF-конструктора, а независимая платформа, которая превращает неструктурированные письма и файлы в проверяемую цифровую модель заявки. Обложка и описание КП являются первым измеримым сценарием, а созданный контур затем может использоваться в подборе замен, технической поддержке, корпоративном поиске и обработке проектной документации.",
    fill=GREEN_FILL,
    label_color=GREEN,
)

add_section_band("I", "СУТЬ РЕШЕНИЯ И ЦЕННОСТЬ ДЛЯ КОМПАНИИ")

h = doc.add_heading("1. Что именно предлагается", level=1)
add_para("Создать локально разворачиваемую AI-платформу, которая принимает почтовую переписку, документы и чертежи, извлекает из них факты, определяет вероятное назначение системы, формирует проект описания и готовит визуальный материал для КП. Решение является самостоятельным: оно содержит собственные механизмы классификации, поиска, ранжирования и проверки и не зависит от каких-либо других нереализованных вариантов автоматизации.")
add_para("Платформа не подменяет инженерное решение. Она собирает и структурирует доступные факты, показывает источники каждого вывода, оценивает уверенность и передаёт человеку случаи, в которых исходных данных недостаточно.")

h = doc.add_heading("2. Почему это стратегический проект, а не генератор текста", level=1)
for b in [
    "Компания получает единый слой работы с неструктурированными данными: письмами, сканами, PDF, таблицами, изображениями и CAD-файлами.",
    "Корпоративные знания перестают оставаться только в переписке и опыте отдельных сотрудников: факты, решения и исправления превращаются в повторно используемые данные.",
    "Каждое подтверждение или исправление менеджера и инженера улучшает качество будущих результатов и формирует собственный датасет компании.",
    "Архитектура допускает локальное размещение моделей и данных, что снижает зависимость от внешних AI-сервисов и позволяет контролировать конфиденциальность.",
    "Созданные сервисы повторно применяются за пределами КП: поиск аналогов, подбор замен, разбор ТЗ, нормализация спецификаций, подготовка ответов и внутренний ассистент.",
]:
    add_bullet(b)

add_callout(
    "Ключевая идея",
    "Компания инвестирует не в одну картинку на первой странице, а в управляемую способность понимать собственный поток технической информации и использовать его в операционных процессах.",
    fill=PALE,
    label_color=BLUE,
)

h = doc.add_heading("3. Рабочий сценарий", level=1)
workflow_rows = [
    ("1. Приём", "Письмо, вложения, выгрузка из CRM или ручная загрузка объединяются в одну заявку."),
    ("2. Разбор", "Текст, таблицы, изображения, реквизиты и элементы чертежей извлекаются в единый JSON-профиль."),
    ("3. Понимание", "Модели определяют сущности, оборудование, ограничения, назначение системы и недостающие сведения."),
    ("4. Подготовка", "Формируются описание, аргументация, черновик визуальной композиции и ссылки на исходные факты."),
    ("5. Контроль", "Верификатор проверяет бренды, модели, характеристики и уверенность; спорные случаи направляются человеку."),
    ("6. Публикация", "Подтверждённые данные передаются в PDF-конструктор и сохраняются как обучающий пример."),
]
add_simple_table(["Этап", "Результат"], workflow_rows, [1900, 7460], font_size=8.9)

h = doc.add_heading("4. Границы автоматизации", level=1)
add_para("Нейросеть не должна реконструировать отсутствующее техническое задание. Если система не находит достаточных оснований для вывода, она обязана сформировать список уточнений и остановить автоматическую публикацию. Такой режим принципиально важен: доверие к платформе создаётся не количеством автоматически обработанных заявок, а контролируемой долей корректных решений.")
boundary_rows = [
    ("Можно автоматизировать", "Извлечение данных, поиск по базе, классификацию, черновики текста, сборку композиции, контроль формальных ошибок."),
    ("Нельзя автоматизировать вслепую", "Выбор назначения при отсутствии признаков, инженерные допущения, подмену неизвестных характеристик, финальное утверждение критичных решений."),
]
add_simple_table(["Контур", "Правило"], boundary_rows, [2500, 6860], font_size=8.9, header_fill=LIGHT)

h = doc.add_heading("5. Практический результат первого контура", level=1)
for b in [
    "структурированная карточка заявки с источниками и уровнем уверенности;",
    "предварительная классификация назначения системы;",
    "редактируемое описание решения без неподтверждённых характеристик;",
    "обложка из реальных фотографий оборудования с управляемой композицией;",
    "список отсутствующих данных и вопросов инженеру;",
    "журнал решения: какие модели, правила и источники повлияли на результат.",
]:
    add_bullet(b)

add_callout(
    "Метрика успеха",
    "На пилоте оцениваются не субъективные впечатления, а точность извлечения, точность классификации, доля заявок без ручного пересбора, время подготовки КП и количество фактических ошибок после проверки.",
    fill=AMBER_FILL,
    label_color=AMBER,
)

doc.add_page_break()
add_section_band("II", "ТЕХНИЧЕСКИЙ ПЛАН РЕАЛИЗАЦИИ")

h = doc.add_heading("6. Архитектурный принцип", level=1)
add_para("Оркестратором не должна быть одна большая модель. Производственный контур строится как граф небольших, наблюдаемых операций. Детерминированные сервисы выполняют преобразование файлов и проверку правил; специализированные модели решают OCR, поиск, классификацию и генерацию; состояние каждой заявки сохраняется между шагами.")

arch_rows = [
    ("1", "Шлюз входящих данных", "Почта/CRM/API, антивирусная проверка, контроль форматов и прав доступа."),
    ("2", "Хранилище исходников", "S3/MinIO для файлов; PostgreSQL для метаданных, статусов и аудита."),
    ("3", "Конвейер извлечения", "OCR, layout-анализ, таблицы, изображения, DWG/DXF и нормализация."),
    ("4", "Каноническая модель заявки", "Единый JSON: объект, контекст, оборудование, ограничения, источники и пробелы."),
    ("5", "RAG и классификация", "Поиск по каталогу/архиву, сопоставление сущностей, назначение и confidence."),
    ("6", "Генераторы", "Текст, объяснение решения, визуальная композиция; реальные фото сохраняются без перерисовки."),
    ("7", "Верификатор и человек", "Проверка фактов, пороги, согласование, корректировка и обратная связь."),
    ("8", "Интеграционный API", "Передача подтверждённой структуры в PDF-конструктор и другие системы."),
]
add_simple_table(["№", "Компонент", "Назначение"], arch_rows, [600, 2500, 6260], font_size=8.2)

h = doc.add_heading("7. Кто оркестрирует модели и сервисы", level=1)
orchestration_rows = [
    ("Kubernetes", "Контейнеры и ресурсы", "Размещение сервисов, GPU/CPU-лимиты, перезапуск, масштабирование. Для пилота допустим Docker Compose."),
    ("Temporal", "Долгий бизнес-процесс", "Очереди, повторные попытки, тайм-ауты, идемпотентность, возобновление заявки после сбоя."),
    ("LangGraph", "AI-граф внутри заявки", "Маршрутизация агентных шагов, состояние, циклы проверки и human-in-the-loop."),
    ("vLLM", "Сервер моделей", "Единый API для LLM/VLM, очередь запросов, батчинг и контроль GPU-памяти."),
    ("FastAPI workers", "Инструменты", "OCR, CAD, каталог, правила, рендеринг, файловые операции и интеграции."),
]
add_simple_table(["Слой", "Что контролирует", "Почему нужен"], orchestration_rows, [1700, 2300, 5360], font_size=8.15)
add_para("Такое разделение исключает ситуацию, в которой языковая модель самостоятельно управляет инфраструктурой. Модель предлагает следующий шаг и структурированный результат; инфраструктурный оркестратор решает, когда и где этот шаг выполнить, а правила безопасности определяют, требуется ли подтверждение человека.", style="Small Note")

h = doc.add_heading("8. Состав специализированных агентов", level=1)
agent_rows = [
    ("Intake Agent", "Собирает письмо и вложения, определяет типы файлов, создаёт паспорт заявки."),
    ("Document Agent", "Запускает OCR/layout, извлекает таблицы и сохраняет координаты источников."),
    ("CAD Agent", "Конвертирует DWG, анализирует слои, блоки, подписи и геометрию."),
    ("Entity Agent", "Нормализует бренды и модели, связывает позиции с каталогом и аналогами."),
    ("Context Agent", "Выделяет назначение, ограничения, зоны, сценарии и противоречия."),
    ("Classification Agent", "Возвращает классы, вероятности, объяснение и признаки, повлиявшие на решение."),
    ("Gap Agent", "Находит недостающие сведения и формирует конкретные вопросы."),
    ("Composer Agent", "Готовит описание системы только на подтверждённых фактах и RAG-контексте."),
    ("Visual Agent", "Собирает коллаж из реальных фото, генерируя только фон, свет и композиционную среду."),
    ("Verifier Agent", "Сверяет модели, бренды, цифры, запрещённые утверждения и порог уверенности."),
]
add_simple_table(["Агент", "Ограниченная ответственность"], agent_rows, [2200, 7160], font_size=8.25)
add_callout(
    "Важно",
    "Под «агентом» понимается не автономный цифровой сотрудник, а ограниченный программный узел с заданной схемой входа/выхода, набором разрешённых инструментов, журналированием и тестами качества.",
    fill=LIGHT,
    label_color=BLUE,
)

h = doc.add_heading("9. Рекомендуемые модели и планирование памяти", level=1)
add_para("Значения ниже являются инженерным резервом для пилотного проектирования, а не паспортными гарантиями. Фактическое потребление зависит от квантизации, длины контекста, размера изображений, числа одновременных запросов и параметров KV-кэша.", style="Small Note")
model_rows = [
    ("Документы", "PaddleOCR-VL 1.6 (0,9B)", "OCR, layout, таблицы", "4-8 ГБ VRAM; 8-16 ГБ RAM"),
    ("Мультимодальный анализ", "Qwen3-VL-8B Instruct", "Письма + страницы + изображения", "16-24 ГБ VRAM; 16-32 ГБ RAM"),
    ("Текст и рассуждение", "Qwen3-14B / 30B-A3B", "Структура, описание, классификация", "18-28 / 28-48 ГБ VRAM в квантизации"),
    ("Корпоративный поиск", "BGE-M3 (569M)", "Мультиязычные embeddings, hybrid retrieval", "2-4 ГБ VRAM или 6-12 ГБ RAM"),
    ("Символы и объекты", "Grounding DINO + дообученный detector", "Обозначения на схемах", "3-8 ГБ VRAM; 8-16 ГБ RAM"),
    ("Визуальная сборка", "ComfyUI + FLUX/Qwen Image Edit", "Фон, свет, тени, композиция", "20-32 ГБ VRAM inference; 24-48 ГБ train"),
    ("CAD", "ODA Converter + ezdxf", "DWG/DXF, слои, блоки, геометрия", "CPU; 4-16 ГБ RAM на задание"),
    ("Сервисы данных", "PostgreSQL/pgvector + MinIO + Redis", "Метаданные, векторы, файлы, кэш", "24-48 ГБ RAM суммарно на пилоте"),
]
add_simple_table(["Задача", "Компонент", "Роль", "Плановый ресурс"], model_rows, [1500, 2600, 2860, 2400], font_size=7.65)

h = doc.add_heading("10. Размещение", level=1)
hardware_rows = [
    ("Пилот", "1× RTX 5090 32 ГБ, 128 ГБ RAM, 4 ТБ NVMe", "Модели запускаются очередями; тяжёлая генерация не работает одновременно с VLM. Подходит для проверки качества и сбора профиля нагрузки."),
    ("Производственная станция", "1× RTX PRO 6000 96 ГБ, 256 ГБ ECC RAM", "Одновременная работа нескольких сервисов, крупные модели, длинный контекст и эксплуатационный запас."),
    ("Масштабирование", "2× RTX PRO 6000, 512 ГБ RAM, отдельное объектное хранилище", "Разделение анализа документов и генерации, параллельные заявки, обучение и отказоустойчивость."),
]
add_simple_table(["Контур", "Конфигурация", "Назначение"], hardware_rows, [1600, 3200, 4560], font_size=8.15)
add_para("Рекомендуемая операционная система: Ubuntu Server LTS, NVIDIA Container Toolkit и контейнеризация. Windows может использоваться на рабочих местах и для отдельных CAD-конвертеров, но центральный AI-контур целесообразно размещать на Linux из-за поддержки CUDA, vLLM и производственных средств мониторинга.", style="Small Note")

h = doc.add_heading("11. Этапы реализации", level=1)
stage_rows = [
    ("0. Проектирование", "2 недели", "Архив данных, классы систем, KPI, безопасность, эталонный набор."),
    ("1. Data foundation", "4-6 недель", "Хранилище, паспорт заявки, коннекторы, OCR и журнал источников."),
    ("2. Понимание", "4-6 недель", "Нормализация оборудования, RAG, классификация, gap-анализ."),
    ("3. Генерация", "6-8 недель", "Тексты, visual pipeline, верификаторы и интеграция с PDF."),
    ("4. Закрытый пилот", "4-6 недель", "Историческая проверка, работа с группой менеджеров, настройка порогов."),
    ("5. Промышленный ввод", "8-12 недель", "Мониторинг, роли, SLA, резервирование, обучение и регламент."),
]
add_simple_table(["Этап", "Срок", "Проверяемый результат"], stage_rows, [2100, 1300, 5960], font_size=8.3)
add_para("Этапы частично выполняются параллельно. Реалистичный горизонт промышленного пилота - 5-7 месяцев; устойчивый производственный контур - 8-12 месяцев при наличии доступа к истории заявок, каталогу и инженерным экспертам.", style="Lead")

h = doc.add_heading("12. Команда и ответственность", level=1)
team_rows = [
    ("Владелец продукта", "1", "Приоритеты, KPI, бюджет, принятие бизнес-результата."),
    ("Solution/AI architect", "1", "Архитектура, безопасность, контракты сервисов, технические решения."),
    ("ML engineers", "2", "VLM/LLM, RAG, классификация, evals, оптимизация inference."),
    ("Backend/data engineers", "2-3", "Коннекторы, хранилища, ETL, API, каталог и аудит."),
    ("Frontend engineer", "1", "Интерфейс проверки, источники, исправления и статусы."),
    ("DevOps/MLOps", "1", "GPU-контур, CI/CD, мониторинг, секреты, резервирование."),
    ("QA/evaluation", "1", "Тестовые наборы, регрессии, качество и приемочные отчёты."),
    ("Инженеры предметной области", "0,5-1 совокупно", "Таксономия, разметка, проверка и правила допустимости."),
    ("Дизайн/контент", "по этапу", "Эталонные композиции, правила бренда и исходные изображения."),
]
add_simple_table(["Роль", "FTE", "Ответственность"], team_rows, [2500, 700, 6160], font_size=8.15)

h = doc.add_heading("13. Контроль качества, безопасность и эксплуатация", level=1)
for b in [
    "Каждый вывод хранит ссылку на фрагмент исходного документа или запись каталога.",
    "Публикация блокируется при низкой уверенности, конфликте источников или отсутствии обязательных полей.",
    "Файлы проходят антивирусную проверку; права на заявки наследуются из корпоративной системы.",
    "Промпты, версии моделей, параметры и ответы журналируются для воспроизводимости и расследования ошибок.",
    "Отдельный evaluation-набор запускается при каждом обновлении модели, правил или справочников.",
    "Персональные и коммерчески чувствительные данные не используются для внешнего обучения без отдельного решения.",
]:
    add_bullet(b)

h = doc.add_heading("14. Управленческий смысл и решение о запуске", level=1)
add_para("Проект создаёт корпоративный AI-фундамент, а не одноразовую автоматизацию. Самая дорогая часть - не закупка GPU и не подключение модели, а перевод накопленных знаний компании в управляемую систему данных, правил, тестов и обратной связи. Именно этот актив обеспечивает долгосрочный эффект и снижает зависимость от конкретного поставщика модели.")
add_callout(
    "Предлагаемое решение",
    "Утвердить двухнедельный этап проектирования: назначить владельца продукта и инженерного владельца таксономии, предоставить обезличенную историческую выборку заявок и итоговых КП, зафиксировать KPI пилота и подготовить смету производственного контура. Покупку максимальной конфигурации и массовое производство визуалов выполнять только после offline-проверки качества.",
    fill=GREEN_FILL,
    label_color=GREEN,
)

h = doc.add_heading("Технические источники", level=1)
sources = [
    "PaddleOCR-VL 1.6: https://www.paddleocr.ai/main/en/version3.x/algorithm/PaddleOCR-VL/PaddleOCR-VL-1.6.html",
    "Qwen3-VL: https://github.com/QwenLM/Qwen3-VL",
    "Qwen3: https://qwenlm.github.io/blog/qwen3/",
    "BGE-M3: https://github.com/FlagOpen/FlagEmbedding/blob/master/docs/source/bge/bge_m3.rst",
    "vLLM supported models: https://docs.vllm.ai/en/latest/models/supported_models/",
    "LangGraph orchestration and human-in-the-loop: https://docs.langchain.com/oss/python/langgraph/overview",
    "ComfyUI: https://github.com/Comfy-Org/ComfyUI",
    "NVIDIA RTX PRO 6000 Blackwell: https://www.nvidia.com/en-us/products/workstations/professional-desktop-gpus/rtx-pro-6000/",
]
for src in sources:
    p = doc.add_paragraph(style="Small Note")
    p.paragraph_format.left_indent = Inches(0.18)
    r = p.add_run(src)
    set_font(r, 7.8, color=MUTED)

# Mark the leading row of every compact grid for screen-reader navigation and
# predictable repetition when a table crosses a page boundary.
for table in doc.tables:
    first_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if first_tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        first_tr_pr.append(tbl_header)

# Keep headings and first content together, prevent widows where practical.
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        set_keep(p, True)
    else:
        set_keep(p, False)

# Core properties: neutral, company-owned artifact.
doc.core_properties.title = "Независимая AI-платформа анализа заявок и формирования КП"
doc.core_properties.subject = "Стратегическая записка и технический план реализации"
doc.core_properties.author = ""
doc.core_properties.keywords = "AI, коммерческое предложение, OCR, RAG, VLM, оркестрация"

doc.save(OUT)
print(OUT)
